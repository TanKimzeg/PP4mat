from typing import Any
from docx.text.paragraph import Paragraph

from pp4mat.logger import setup_logger
from pp4mat.config_converter import FormatConfig
from pp4mat.format_checker.protocols import FormatChecker, FormatErrors

logger = setup_logger(__package__)

class TOCChecker(FormatChecker):
    def __init__(self, sections: dict[str, list[Paragraph]], doc: Any, config: FormatConfig, errors: FormatErrors) -> None:
        self.errors = errors
        self.paragraphs = sections["目录"]
        self.doc = doc
        self.config = config
        win32doc = self.doc
        toc_paragraphs: list[Paragraph] = self.paragraphs
        # 目录兜底：部分 docx 的自动目录可能不在 python-docx 的 paragraphs 视图里。
        # 用 COM 直接读取 TablesOfContents 的 Range.Text。
        try:
            if not toc_paragraphs and win32doc.TablesOfContents.Count >= 1:
                toc_text = win32doc.TablesOfContents(1).Range.Text or ""
                toc_lines = [ln.strip() for ln in str(toc_text).splitlines() if ln.strip()]

                class _FakeParagraph:
                    def __init__(self, text: str):
                        self.text = text

                # 仅 toc_checker 用到 p.text，所以用轻量对象即可
                toc_paragraphs = [
                    _FakeParagraph(t)  # type: ignore[arg-type]
                    for t in toc_lines
                ]
                logger.info(f"已使用 Word COM 提取目录，共 {len(toc_paragraphs)} 行")
        except Exception as e:
            logger.warning(f"Word COM 提取目录失败，仍使用 python-docx 目录段落：{e}")
        finally:
            self.paragraphs = toc_paragraphs
            sections["目录"] = self.paragraphs


    def check_format(self) -> None:
        self.toc_checker(self.paragraphs, self.config, self.errors)
        return None

    @staticmethod
    def toc_checker(toc: list[Paragraph], format_config: FormatConfig, errors: dict[str, list[str]]) -> None:
        logger.info("检查目录格式...")
        if not toc:
            logger.warning("目录部分为空，跳过检查。")
            return
        import re
        h1_pattern = re.compile(r"^\d+\.")
        h2_pattern = re.compile(r"^\d+\.\d+")
        h3_pattern = re.compile(r"^\d+\.\d+\.\d+")
        chapter_count = 0
        for p in toc:
            if re.match(h1_pattern, p.text.strip()):
                if not re.match(h2_pattern, p.text.strip()):
                    chapter_count += 1
                if re.match(h3_pattern, p.text.strip()):
                    logger.error(f"目录项格式错误：{p.text.strip()}，不应包含三级标题")
                    errors["目录检测"].append(f"目录项格式错误：{p.text.strip()}，不应包含三级标题")
        if format_config.chapter_min_count and chapter_count < format_config.chapter_min_count:
            logger.error(f"目录章节数量少于{format_config.chapter_min_count}章，当前数量为{chapter_count}章，请检查！")
            errors["目录检测"].append(f"目录章节数量少于{format_config.chapter_min_count}章，当前数量为{chapter_count}章")
